#!/usr/bin/env python3
"""Soluzione completa di un esercizio di progettazione da un modello E/R (file .er), in un file solo.

Uso:
    python3 soluzione.py MODELLO.er [CARTELLA] [--file=html|docx|md|html,docx]
                         [--parti=traccia,ipotesi,diagramma,documentazione,schema,query] [--scala=1.25]

Scrive in CARTELLA (predefinita: quella del file .er):
  - MODELLO.html (predefinito): una pagina sola, con il diagramma dentro il file; si apre nel browser,
    si stampa o si salva in PDF, si carica su Moodle o Classroom, si incolla in un Google Doc;
  - MODELLO.docx (--file=docx): documento Word, modificabile, con il diagramma dentro (serve python-docx);
  - MODELLO.md (--file=md): Markdown per il libro, con il diagramma come file a parte (MODELLO.png);
  - sempre, se c'è il diagramma: MODELLO.drawio e MODELLO.png.

Parti, in quest'ordine (predefinite tutte; quelle vuote si omettono): traccia (blocco «traccia:» della
descrizione), ipotesi aggiuntive, diagramma E/R (servono le posizioni), documentazione (entità, associazioni,
vincoli non esprimibili), schema relazionale (con i vincoli), query SQL (blocchi «query:»). Ogni parte si può produrre anche da sola con il
suo script: disegna_er.py, documentazione.py, schema_relazionale.py.
"""
import base64
import html
import os
import re
import sys

QUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, QUI)
from modello_er import leggi  # noqa: E402
import documentazione as doc  # noqa: E402
import schema_relazionale as sr  # noqa: E402

PARTI = ("traccia", "ipotesi", "diagramma", "documentazione", "schema", "query")
FORMATI = ("html", "docx", "md")


# ---------------------------------------------------------------- contenuto comune ai formati

def prepara(testo_er, parti, base, scala):
    """Blocchi: ("h1"|"h2"|"h3", testo), ("p", testo), ("ul", [testi]), ("img", percorso, alt),
    ("tabella", intestazione, righe), ("schema", relazioni, vincoli)."""
    m = leggi(testo_er)
    b = [("h1", m["titolo"])]
    if "traccia" in parti and m["traccia"]:
        b.append(("h2", "Traccia"))
        b += [("p", t) for t in m["traccia"]]
    if "ipotesi" in parti and m["ipotesi"]:
        b += [("h2", "Ipotesi aggiuntive"), ("ul", list(m["ipotesi"]))]
    if "diagramma" in parti:
        import disegna_er
        from drawio_png import disegna, leggi_modello
        open(base + ".drawio", "w", encoding="utf-8").write(disegna_er.costruisci(leggi(testo_er)))
        disegna(leggi_modello(base + ".drawio"), scala).save(base + ".png")
        print("OK: %s.drawio e %s.png" % (base, base))
        b += [("h2", "Diagramma E/R"), ("img", base + ".png", "Diagramma E/R: " + m["titolo"])]
    if "documentazione" in parti:
        t, avvisi = doc.tabelle(leggi(testo_er))
        for av in avvisi:
            print("AVVISO: " + av)
        b += [("h2", "Documentazione"), ("h3", "Documentazione delle entità"),
              ("tabella", doc.INTEST_ENTITA, t["entita"])]
        if t["associazioni"]:
            b += [("h3", "Documentazione delle associazioni"), ("tabella", doc.INTEST_ASSOC, t["associazioni"])]
        if m["vincoli"]:
            b += [("h3", "Vincoli non esprimibili"), ("ul", list(m["vincoli"]))]
    if "schema" in parti:
        relazioni, vincoli = sr.traduci(leggi(testo_er))
        for nota in sr.NOTE:
            print("NOTA: " + nota)
        for av in sr.AVVISI:
            print("AVVISO: " + av)
        b += [("h2", "Schema relazionale"), ("schema", relazioni, vincoli)]
    if "query" in parti and m["query"]:
        b.append(("h2", "Query SQL"))
        for i, q in enumerate(m["query"], 1):
            b += [("h3", "%d. %s" % (i, q["titolo"])), ("codice", "\n".join(q["sql"]).strip("\n"))]
    return b


def pezzi_inline(testo):
    """**grassetto** e *corsivo* -> [(testo, grassetto, corsivo)]."""
    out = []
    for parte in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)", testo):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            out.append((parte[2:-2], True, False))
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            out.append((parte[1:-1], False, True))
        else:
            out.append((parte, False, False))
    return out


def gruppi(righe):
    """Righe raggruppate per entità o associazione (la prima cella in grassetto apre un gruppo)."""
    out = []
    for r in righe:
        if isinstance(r[0], tuple) or not out:
            out.append([r])
        else:
            out[-1].append(r)
    return out


FISSE = {6: 3, 7: 4}            # colonne che si uniscono in verticale: entità 3, associazioni 4
LARGHEZZE = {6: [2.4, 2.0, 3.4, 3.5, 2.3, 3.4], 7: [2.4, 2.6, 2.4, 3.2, 2.4, 2.0, 2.0]}   # cm, somma 17


# ---------------------------------------------------------------- HTML (un file solo)

# Gli stili stanno negli attributi, non in un <style>: Moodle toglie i <style>, e la pagina si può incollare
# nell'editor (codice sorgente) senza perdere l'aspetto. Lo stesso aspetto lo dà la dashboard «Tracce di
# progettazione» (uscite_traccia.js di tracce-progettazione, CSS_SOL): se cambia qui, va cambiato anche lì.
STILI = {
    "box": "font-family:Calibri,'Segoe UI',Helvetica,Arial,sans-serif;color:#1a1a1a;line-height:1.5;max-width:900px",
    "h1": "font-size:1.6em;font-weight:normal;color:#1c4587;border-bottom:2px solid #6c8ebf;padding-bottom:6px;margin:0 0 .8em",
    "h2": "font-size:1.25em;font-weight:bold;color:#1c4587;margin:2em 0 .4em",
    "h3": "font-size:1.05em;font-weight:bold;color:#1c4587;margin:1.4em 0 .4em",
    "p": "margin:0 0 .8em", "ul": "margin:0 0 .8em;padding-left:1.6em",
    "img": "max-width:100%;height:auto;display:block;margin:12px 0",
    "table": "border-collapse:collapse;width:100%;font-size:.9em;margin:8px 0 16px",
    "th": "border:1px solid #999999;padding:4px 8px;text-align:left;vertical-align:top;background:#eeeeee",
    "td": "border:1px solid #999999;padding:4px 8px;text-align:left;vertical-align:top",
    "inizio": "border-top:2px solid #666666",
    "relazione": "margin:6px 0;font-size:1.05em",
    "pre": "background:#f5f7fb;border:1px solid #c9d3e3;border-radius:4px;padding:10px 12px;overflow-x:auto;"
           "font-family:Consolas,'Courier New',monospace;font-size:.9em;line-height:1.4;margin:6px 0 12px",
}


def st(k):
    return ' style="%s"' % STILI[k]


def inline_html(testo):
    out = ""
    for t, gr, co in pezzi_inline(testo):
        t = html.escape(t)
        out += "<strong>%s</strong>" % t if gr else ("<em>%s</em>" % t if co else t)
    return out


def rendi_html(blocchi):
    corpo = []
    titolo = "Soluzione"
    for bl in blocchi:
        k = bl[0]
        if k in ("h1", "h2", "h3"):
            if k == "h1":
                titolo = bl[1]
            corpo.append("<%s%s>%s</%s>" % (k, st(k), html.escape(bl[1]), k))
        elif k == "p":
            corpo.append("<p%s>%s</p>" % (st("p"), inline_html(bl[1])))
        elif k == "ul":
            corpo.append("<ul%s>\n%s\n</ul>" % (st("ul"), "\n".join("<li>%s</li>" % inline_html(x) for x in bl[1])))
        elif k == "img":
            dati = base64.b64encode(open(bl[1], "rb").read()).decode("ascii")
            corpo.append('<img src="data:image/png;base64,%s" alt="%s"%s>' % (dati, html.escape(bl[2], quote=True), st("img")))
        elif k == "tabella":
            intest, righe = bl[1], bl[2]
            fisse = FISSE.get(len(intest), 0)
            r_html = []
            for g in gruppi(righe):
                for j, r in enumerate(g):
                    celle = []
                    for i, c in enumerate(r):
                        if i < fisse and j > 0:
                            continue
                        span = ' rowspan="%d"' % len(g) if i < fisse and len(g) > 1 else ""
                        testo = "<strong>%s</strong>" % html.escape(c[0]) if isinstance(c, tuple) else html.escape(c)
                        stile = STILI["td"] + (";" + STILI["inizio"] if j == 0 else "")
                        celle.append('<td%s style="%s">%s</td>' % (span, stile, testo))
                    r_html.append("<tr>%s</tr>" % "".join(celle))
            corpo.append("<table%s>\n<thead><tr>%s</tr></thead>\n<tbody>\n%s\n</tbody>\n</table>"
                         % (st("table"), "".join("<th%s>%s</th>" % (st("th"), html.escape(h)) for h in intest), "\n".join(r_html)))
        elif k == "codice":
            corpo.append("<pre%s>%s</pre>" % (st("pre"), html.escape(bl[1])))
        elif k == "schema":
            relazioni, vincoli = bl[1], bl[2]
            for r in relazioni:
                campi = []
                for c in sr.campi_ordinati(r):
                    t = html.escape(c.nome)
                    if c.pk:
                        t = "<u>%s</u>" % t
                    campi.append(t + ("*" if c.fk else ""))
                corpo.append('<p%s><strong>%s</strong>(%s)</p>' % (st("relazione"), html.escape(r.nome), ", ".join(campi)))
            if vincoli:
                corpo.append("<h3%s>Vincoli</h3>\n<ul%s>\n%s\n</ul>" % (st("h3"), st("ul"), "\n".join("<li>%s</li>" % html.escape(v) for v in vincoli)))
    return ('<!doctype html>\n<html lang="it">\n<head>\n<meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width, initial-scale=1">\n<title>%s</title>\n'
            '</head>\n<body style="background:#ffffff;margin:32px auto;padding:0 20px;max-width:900px">\n'
            '<div%s>\n%s\n</div>\n</body>\n</html>\n' % (html.escape(titolo), st("box"), "\n".join(corpo)))


# ---------------------------------------------------------------- Markdown (per il libro)

def rendi_md(blocchi, base_nome):
    out = []
    for bl in blocchi:
        k = bl[0]
        if k in ("h1", "h2", "h3"):
            out.append("#" * int(k[1]) + " " + bl[1])
        elif k == "p":
            out.append(bl[1])
        elif k == "ul":
            out.append("\n".join("- " + x for x in bl[1]))
        elif k == "img":
            out.append("![%s](%s.png)" % (bl[2], base_nome))
        elif k == "tabella":
            out.append(doc.tabella_md(bl[1], bl[2]))
        elif k == "codice":
            out.append("```sql\n%s\n```" % bl[1])
        elif k == "schema":
            out.append(sr.formatta(bl[1], bl[2], "libro").strip())
    return "\n\n".join(out) + "\n"


# ---------------------------------------------------------------- Word

def rendi_docx(blocchi, percorso):
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError:
        raise SystemExit("ERRORE: per il file Word serve python-docx (pip install python-docx)")
    from PIL import Image

    d = Document()
    sez = d.sections[0]
    sez.page_width, sez.page_height = Cm(21), Cm(29.7)
    for lato in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sez, lato, Cm(2))
    stile = d.styles["Normal"]
    stile.font.name = "Calibri"
    stile.font.size = Pt(11)
    larghezza_utile = Cm(17)

    def ombra(cella, colore="EEEEEE"):
        tcPr = cella._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), colore)
        tcPr.append(shd)

    def paragrafo_inline(p, testo):
        for t, gr, co in pezzi_inline(testo):
            run = p.add_run(t)
            run.bold, run.italic = gr, co

    for bl in blocchi:
        k = bl[0]
        if k in ("h1", "h2", "h3"):
            h = d.add_heading(bl[1], level={"h1": 0, "h2": 1, "h3": 2}[k])
            h.paragraph_format.keep_with_next = True
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x1C, 0x45, 0x87)
        elif k == "p":
            paragrafo_inline(d.add_paragraph(), bl[1])
        elif k == "ul":
            for x in bl[1]:
                paragrafo_inline(d.add_paragraph(style="List Bullet"), x)
        elif k == "img":
            with Image.open(bl[1]) as im:
                larg_cm = im.width / 96 * 2.54
            d.add_picture(bl[1], width=min(larghezza_utile, Cm(larg_cm)))
        elif k == "tabella":
            intest, righe = bl[1], bl[2]
            n = len(intest)
            fisse = FISSE.get(n, 0)
            larg = [Cm(x) for x in LARGHEZZE.get(n, [17 / n] * n)]
            t = d.add_table(rows=1, cols=n)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            for i, col in enumerate(t.columns):
                col.width = larg[i]
            intestazione = t.rows[0]
            # la riga d'intestazione si ripete a ogni pagina e resta con la riga che segue
            trPr = intestazione._tr.get_or_add_trPr()
            el = OxmlElement("w:tblHeader")
            el.set(qn("w:val"), "true")
            trPr.append(el)
            for i, h in enumerate(intest):
                c = intestazione.cells[i]
                c.width = larg[i]
                par = c.paragraphs[0]
                par.paragraph_format.keep_with_next = True
                run = par.add_run(h)
                run.bold = True
                run.font.size = Pt(9)
                ombra(c)
            for g in gruppi(righe):
                prime = None
                for j, r in enumerate(g):
                    riga = t.add_row()
                    riga._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                    celle = riga.cells
                    for i, x in enumerate(r):
                        celle[i].width = larg[i]
                        if i < fisse and j > 0:
                            continue
                        testo, gr = (x[0], True) if isinstance(x, tuple) else (x, False)
                        par = celle[i].paragraphs[0]
                        if j < len(g) - 1:
                            par.paragraph_format.keep_with_next = True
                        run = par.add_run(testo)
                        run.bold = gr
                        run.font.size = Pt(9)
                    if j == 0:
                        prime = celle
                    elif fisse:
                        for i in range(fisse):
                            prime[i].merge(celle[i])
            d.add_paragraph()
        elif k == "codice":
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            for i, riga in enumerate(bl[1].split("\n")):
                if i:
                    p.add_run().add_break()
                run = p.add_run(riga)
                run.font.name = "Consolas"
                run.font.size = Pt(10)
        elif k == "schema":
            relazioni, vincoli = bl[1], bl[2]
            for r in relazioni:
                p = d.add_paragraph()
                p.add_run(r.nome).bold = True
                p.add_run("(")
                campi = sr.campi_ordinati(r)
                for i, c in enumerate(campi):
                    run = p.add_run(c.nome)
                    run.underline = c.pk
                    if c.fk:
                        p.add_run("*")
                    if i < len(campi) - 1:
                        p.add_run(", ")
                p.add_run(")")
            if vincoli:
                d.add_heading("Vincoli", level=2)
                for v in vincoli:
                    d.add_paragraph(v, style="List Bullet")
    d.save(percorso)


# ---------------------------------------------------------------- main

def main():
    a = [x for x in sys.argv[1:] if not x.startswith("--")]
    if not a:
        raise SystemExit(__doc__)
    opz = dict(x[2:].split("=", 1) if "=" in x else (x[2:], "1") for x in sys.argv[1:] if x.startswith("--"))
    parti = [p.strip() for p in opz.get("parti", ",".join(PARTI)).split(",") if p.strip()]
    formati = [f.strip() for f in opz.get("file", "html").split(",") if f.strip()]
    for p in parti:
        if p not in PARTI:
            raise SystemExit("ERRORE: parte «%s» sconosciuta (parti: %s)" % (p, ", ".join(PARTI)))
    for f in formati:
        if f not in FORMATI:
            raise SystemExit("ERRORE: formato «%s» sconosciuto (formati: %s)" % (f, ", ".join(FORMATI)))
    src = a[0]
    cartella = a[1] if len(a) > 1 else os.path.dirname(os.path.abspath(src))
    os.makedirs(cartella, exist_ok=True)
    base_nome = os.path.splitext(os.path.basename(src))[0]
    base = os.path.join(cartella, base_nome)
    blocchi = prepara(open(src, encoding="utf-8").read(), parti, base, float(opz.get("scala", 1.25)))
    if "html" in formati:
        open(base + ".html", "w", encoding="utf-8").write(rendi_html(blocchi))
        print("OK: %s.html" % base)
    if "docx" in formati:
        rendi_docx(blocchi, base + ".docx")
        print("OK: %s.docx" % base)
    if "md" in formati:
        open(base + ".md", "w", encoding="utf-8").write(rendi_md(blocchi, base_nome))
        print("OK: %s.md" % base)


if __name__ == "__main__":
    main()
